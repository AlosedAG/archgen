"""Shared exporters for the note-taking/tracking modules — Written
Requirements Document, Joint Evaluation Plan, Test Case Document — plus
anything else that just needs "editable tables out as a document."

These modules are structured note-taking, not generated output like the
Architecture Generator or Portal Auditor: there's no computation to test,
just consistent formatting across four downloadable formats. A ``Section``
is a ``(title, rows)`` pair, where each row is an ordered dict of
``column -> value``; column order is taken from the first row's keys, so
callers should build rows with consistent key order (e.g. from a
``pandas.DataFrame.to_dict("records")``).

Every format here is picked to require no system binary: python-docx and
openpyxl write their formats directly, and fpdf2 is a pure-Python PDF
writer — none of them shell out to Word, LibreOffice, or a native PDF
engine, so this works the same on any machine that can run the app.
"""

from __future__ import annotations

import csv
import io

import pandas as pd
from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

Section = tuple[str, list[dict[str, str]]]


def _columns(rows: list[dict[str, str]]) -> list[str]:
    return list(rows[0].keys()) if rows else []


def clean_rows(df: pd.DataFrame, required_columns: list[str] | None = None) -> list[dict[str, str]]:
    """Convert an edited ``st.data_editor`` DataFrame into clean
    string-valued rows for a :data:`Section`, dropping rows where every
    column in ``required_columns`` (default: all columns) is blank —
    covers the trailing empty row Streamlit's dynamic-rows editor always
    offers, plus any row a user added and then abandoned."""
    rows: list[dict[str, str]] = []
    for record in df.to_dict("records"):
        cleaned = {col: "" if pd.isna(value) else str(value) for col, value in record.items()}
        required = required_columns or list(cleaned.keys())
        if any(cleaned.get(col, "").strip() for col in required):
            rows.append(cleaned)
    return rows


def sections_to_docx(title: str, subtitle: str, sections: list[Section]) -> io.BytesIO:
    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        doc.add_paragraph(subtitle)
    for section_title, rows in sections:
        doc.add_heading(section_title, level=1)
        columns = _columns(rows)
        if not columns:
            doc.add_paragraph("None recorded.")
            continue
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        for cell, header in zip(table.rows[0].cells, columns):
            cell.text = header
        for row in rows:
            cells = table.add_row().cells
            for cell, col in zip(cells, columns):
                cell.text = str(row.get(col, "") or "")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def sections_to_xlsx(sections: list[Section]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    for section_title, rows in sections:
        name = (section_title or "Sheet")[:31]
        base, suffix_i = name, 1
        while name in used_names:
            suffix = f" ({suffix_i})"
            name = base[: 31 - len(suffix)] + suffix
            suffix_i += 1
        used_names.add(name)
        ws = wb.create_sheet(title=name)

        columns = _columns(rows)
        if not columns:
            ws.append(["None recorded."])
            continue
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(col, "") for col in columns])
        for idx, col in enumerate(columns, start=1):
            width = max([len(col)] + [len(str(row.get(col, "") or "")) for row in rows]) + 2
            ws.column_dimensions[get_column_letter(idx)].width = min(width, 60)

    if not wb.sheetnames:
        wb.create_sheet(title="Sheet1")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sections_to_csv(sections: list[Section]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for i, (section_title, rows) in enumerate(sections):
        if i > 0:
            writer.writerow([])
        writer.writerow([f"## {section_title}"])
        columns = _columns(rows)
        if not columns:
            writer.writerow(["None recorded."])
            continue
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(col, "") for col in columns])
    return buffer.getvalue()


# fpdf2's built-in core fonts (Helvetica etc.) only support Latin-1, but
# this app's own copy — dashes, curly quotes, arrows, middots — leans on
# Unicode punctuation throughout. Rather than bundle a Unicode TTF just
# for the PDF export, swap the common cases for ASCII equivalents and
# replace anything still outside Latin-1 (e.g. emoji in free-text fields)
# rather than letting fpdf2 raise.
_PDF_SAFE_REPLACEMENTS = {
    "—": "-",
    "–": "-",
    "‘": "'",
    "’": "'",
    "“": '"',
    "”": '"',
    "•": "-",
    "·": "-",
    "→": "->",
    "…": "...",
}


def _pdf_safe(text: str) -> str:
    for src, dst in _PDF_SAFE_REPLACEMENTS.items():
        text = text.replace(src, dst)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def sections_to_pdf(title: str, subtitle: str, sections: list[Section]) -> bytes:
    pdf = FPDF(orientation="L")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # multi_cell defaults to leaving the cursor at the end of the last
    # line rather than back at the left margin, which starves the next
    # multi_cell of horizontal space -- reset explicitly after each call.
    def _block(text: str, height: float) -> None:
        pdf.multi_cell(0, height, _pdf_safe(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font("Helvetica", "B", 16)
    _block(title, 10)
    if subtitle:
        pdf.set_font("Helvetica", "", 10)
        _block(subtitle, 6)
    pdf.ln(2)

    for section_title, rows in sections:
        pdf.set_font("Helvetica", "B", 13)
        _block(section_title, 8)
        pdf.set_font("Helvetica", "", 9)

        columns = _columns(rows)
        if not columns:
            _block("None recorded.", 6)
            pdf.ln(3)
            continue

        with pdf.table(text_align="LEFT") as table:
            header_row = table.row()
            for col in columns:
                header_row.cell(_pdf_safe(col))
            for row in rows:
                data_row = table.row()
                for col in columns:
                    data_row.cell(_pdf_safe(str(row.get(col, "") or "")))
        pdf.ln(3)

    return bytes(pdf.output())
