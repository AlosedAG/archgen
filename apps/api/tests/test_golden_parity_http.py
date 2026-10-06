"""Golden parity over HTTP: for inputs identical to the golden harness's,
the API returns exactly what the Streamlit-era code produced."""

from __future__ import annotations

import io
import json
from collections.abc import Callable
from pathlib import Path
from typing import Any

from docx import Document
from httpx import AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader

Golden = Callable[[str], Any]


def _without(d: dict[str, Any], *keys: str) -> dict[str, Any]:
    return {k: v for k, v in d.items() if k not in keys}


def _docx_lines(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    lines += [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
    return lines


def _pdf_pages(data: bytes) -> list[str]:
    return [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]


def _name_counts_to_pairs(context: dict[str, Any]) -> dict[str, Any]:
    """API exposes (name, count) pairs as objects; golden fixtures hold the
    domain's tuples. Convert back for comparison."""
    ctx = json.loads(json.dumps(context))
    for s in ctx["object_summaries"]:
        s["top_groups"] = [[g["name"], g["count"]] for g in s["top_groups"]]
    for p in ctx["pipeline_summaries"]:
        p["pipelines"] = [[g["name"], g["count"]] for g in p["pipelines"]]
    return ctx


async def _sample_blueprint(client: AsyncClient, repo_root: Path) -> dict[str, Any]:
    """The blueprint for the sample project, as a client receives it (key
    order preserved — golden fixtures are stored key-sorted, so they can't
    be posted back where dict order drives output order)."""
    sample = json.loads((repo_root / "sample_inputs" / "example_project.json").read_text("utf-8"))
    response = await client.post("/api/v1/blueprints", json=sample)
    assert response.status_code == 200, response.text
    blueprint: dict[str, Any] = response.json()
    return blueprint


async def test_blueprint_matches_golden(client: AsyncClient, golden: Golden, repo_root: Path) -> None:
    expected = golden("blueprint")["blueprint"]
    assert _without(await _sample_blueprint(client, repo_root), "generated_at") == _without(expected, "generated_at")


async def test_blueprint_markdown_matches_golden(client: AsyncClient, golden: Golden, repo_root: Path) -> None:
    blueprint = await _sample_blueprint(client, repo_root)
    response = await client.post("/api/v1/blueprints/markdown", json=blueprint)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert response.text.replace(blueprint["generated_at"], "<TIMESTAMP>") == golden("blueprint")["markdown"]


async def test_diagram_from_blueprint_matches_golden(client: AsyncClient, golden: Golden, repo_root: Path) -> None:
    response = await client.post(
        "/api/v1/diagrams/from-blueprint", json={"blueprint": await _sample_blueprint(client, repo_root)}
    )
    assert response.status_code == 200, response.text
    body = response.json()
    expected = golden("diagram")
    assert {"nodes": body["nodes"], "edges": body["edges"]} == expected["from_blueprint"]
    assert body["dangling_edges"] == expected["dangling"]


async def test_diagram_dot_and_drawio_render_match_golden(client: AsyncClient, golden: Golden, repo_root: Path) -> None:
    graph = (
        await client.post("/api/v1/diagrams/from-blueprint", json={"blueprint": await _sample_blueprint(client, repo_root)})
    ).json()
    dot = await client.post("/api/v1/diagrams/render?format=dot", json={"diagram": graph})
    drawio = await client.post("/api/v1/diagrams/render?format=drawio", json={"diagram": graph})
    assert dot.text.splitlines() == golden("diagram")["dot"]
    assert drawio.text.splitlines() == golden("diagram")["drawio"]
    assert dot.headers["content-disposition"] == 'attachment; filename="architecture_diagram.dot"'


async def test_report_context_matches_golden(client: AsyncClient, golden: Golden) -> None:
    ctx = golden("report")["context"]
    findings = [item for group in ctx["finding_groups"] for item in group["items"]]
    response = await client.post(
        "/api/v1/reports/context", json={"snapshot": ctx["snapshot"], "findings": findings, "project_name": "Golden"}
    )
    assert response.status_code == 200, response.text
    assert _without(_name_counts_to_pairs(response.json()), "generated_at") == _without(ctx, "generated_at")


async def test_unaudited_report_context_matches_golden(client: AsyncClient, golden: Golden) -> None:
    ctx = golden("report")["context_unaudited"]
    response = await client.post(
        "/api/v1/reports/context", json={"snapshot": ctx["snapshot"], "findings": None, "project_name": "Golden"}
    )
    assert _without(_name_counts_to_pairs(response.json()), "generated_at") == _without(ctx, "generated_at")


async def test_discovery_payload_matches_golden(client: AsyncClient, golden: Golden) -> None:
    """Same notes as the golden harness (its one legacy free-text answer is
    sent as structured notes, which flattens to the same line)."""
    answer = {"selected": [], "other": "", "notes": "", "flagged": False}
    notes = {
        "client_name": "Bright Smiles",
        "business_type": "dental",
        "sections": {
            "Goals": {
                "answers": {
                    "kpis": {
                        "selected": ["Win rate", "Other (specify)"],
                        "other": "demo-to-close",
                        "notes": "25%",
                        "flagged": False,
                    },
                    "timeline": {**answer, "selected": ["1–3 months"], "flagged": True},
                    "long_term_goals": {**answer, "notes": "open 3 clinics"},
                }
            },
            "Data": {
                "answers": {"data_regions": {**answer, "selected": ["EU / UK", "United States"]}},
                "edge_cases": "Two legal entities",
            },
            "Processes": {"answers": {"pain_points": {**answer, "notes": "legacy free text answer"}}},
        },
    }
    response = await client.post("/api/v1/discovery/business-analysis/payload", json=notes)
    assert response.status_code == 200, response.text
    assert response.json() == golden("discovery")["business_analysis"]


async def test_suggest_modules_matches_golden(client: AsyncClient, golden: Golden) -> None:
    md = "## Recommended HubSpot Modules\n- **Sales Hub** — pipelines\n- Ops hub: ERP sync\n## Open Questions & Gaps\n- Marketing"
    response = await client.post("/api/v1/discovery/suggest-modules", json={"business_analysis_markdown": md})
    assert response.json() == {"modules": golden("discovery")["suggest_modules"]}


_SECTIONS = {
    "title": "Golden WRD",
    "subtitle": "Sub",
    "sections": [
        {
            "title": "Customer Requirements",
            "rows": [{"Requirement": "Sync contacts nightly"}, {"Requirement": "Log SMS replies"}],
        },
        {"title": "Open Questions", "rows": []},
    ],
}
_MARKDOWN = "# Title\n## Section — one\n- bullet → arrow · middot\n**bold** text\n### Sub\nPlain “quoted” paragraph."


async def test_section_exports_match_golden(client: AsyncClient, golden: Golden) -> None:
    expected = golden("exporters")
    docx = await client.post("/api/v1/exports/sections?format=docx", json=_SECTIONS)
    xlsx = await client.post("/api/v1/exports/sections?format=xlsx", json=_SECTIONS)
    csv = await client.post("/api/v1/exports/sections?format=csv", json=_SECTIONS)
    pdf = await client.post("/api/v1/exports/sections?format=pdf", json=_SECTIONS)
    assert _docx_lines(docx.content) == expected["docx"]
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in workbook.worksheets} == expected["xlsx"]
    assert csv.text.splitlines() == expected["csv"]
    assert _pdf_pages(pdf.content) == expected["pdf"]
    assert pdf.headers["content-type"] == "application/pdf"


async def test_markdown_exports_match_golden(client: AsyncClient, golden: Golden) -> None:
    expected = golden("exporters")
    docx = await client.post("/api/v1/exports/markdown?format=docx", json={"markdown": _MARKDOWN, "title": "Doc"})
    pdf = await client.post("/api/v1/exports/markdown?format=pdf", json={"markdown": _MARKDOWN, "title": "Doc"})
    assert _docx_lines(docx.content) == expected["md_docx"]
    assert _pdf_pages(pdf.content) == expected["md_pdf"]


def test_golden_fixtures_are_the_domain_baseline(repo_root: Path) -> None:
    """Guard: these tests read the domain package's golden fixtures, not a copy."""
    assert (repo_root / "packages" / "domain" / "tests" / "golden" / "fixtures" / "blueprint.json").is_file()
    assert Path(__file__).parent.name == "tests"
