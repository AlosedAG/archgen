"""Deal Stage Planner endpoints: the pre-call template downloads, a
filled-out copy parses into stages plus a preview of what it populates,
and non-planner uploads get a clear problem response."""

from __future__ import annotations

import io

from httpx import AsyncClient
from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def test_template_downloads_as_xlsx(client: AsyncClient) -> None:
    response = await client.get("/api/v1/stage-planner/template")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert 'filename="HubSpot_Deal_Stage_Planner.xlsx"' in response.headers["content-disposition"]
    assert load_workbook(io.BytesIO(response.content)).sheetnames == ["Instructions", "Planner"]


async def test_import_round_trips_the_template(client: AsyncClient) -> None:
    template = (await client.get("/api/v1/stage-planner/template")).content
    response = await client.post("/api/v1/stage-planner/import", content=template, headers={"Content-Type": XLSX})
    assert response.status_code == 200
    body = response.json()
    assert [s["name"] for s in body["stages"]][:2] == ["Appointment Scheduled", "Qualified to Buy"]
    assert body["warnings"] == []
    assert body["pipeline"]["stages"][0]["probability"] == 0.1
    assert body["discovery_answer"].startswith("Appointment Scheduled (10%) → Qualified to Buy (20%)")
    close = next(p for p in body["properties"] if p["label"] == "Expected Close Date")
    assert close["existing"] == "closedate"
    assert close["definition"] is None
    assert any(w["trigger"] == "Deal enters the 'Contract Sent' stage" for w in body["workflows"])
    assert any(i["section"] == "pipelines" for i in body["build_items"])


async def test_import_rejects_a_non_workbook(client: AsyncClient) -> None:
    response = await client.post("/api/v1/stage-planner/import", content=b"not a spreadsheet")
    assert response.status_code == 422
    assert response.json()["type"].endswith("/not-a-stage-planner")


async def test_preview_maps_edited_stages_and_evaluate_adds_their_build_items(client: AsyncClient) -> None:
    stages = [
        {"name": "Lead", "forecast": 10, "automations": "Notify owner", "fields": ["Budget"]},
        {"name": "Won", "forecast": 100},
    ]
    body = (await client.post("/api/v1/stage-planner/preview", json={"stages": stages})).json()
    assert [s["name"] for s in body["pipeline"]["stages"]] == ["Lead", "Won"]
    assert body["pipeline"]["stages"][1]["is_won"] is True
    assert body["discovery_notes"].startswith("[Deal Stage Planner]")
    assert body["workflows"][0]["trigger"] == "Deal enters the 'Lead' stage"

    evaluated = (await client.post("/api/v1/discovery/evaluate", json={"stage_planner": stages})).json()
    pipelines = next(s for s in evaluated["build_order"]["sections"] if s["id"] == "pipelines")
    assert pipelines["items"][0]["name"] == "Deal pipeline from the Stage Planner (2 stages)"
