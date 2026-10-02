"""Shared exporters for the note-taking/tracking modules — Written
Requirements Document, Joint Evaluation Plan, Test Case Document — plus
anything else that just needs "editable tables out as a document."

These modules are structured note-taking, not generated output like the
Architecture Generator or Portal Auditor: there's no computation to test,
just consistent formatting across four downloadable formats. A ``Section``
is a ``(title, rows)`` pair, where each row is an ordered dict of
``column -> value``; column order is taken from the first row's keys, so
callers should build rows with consistent key order (e.g. from a
``pandas.DataFrame.to_dict("records")``).

Every format here is picked to require no system binary: python-docx and
openpyxl write their formats directly, and fpdf2 is a pure-Python PDF
writer — none of them shell out to Word, LibreOffice, or a native PDF
engine, so this works the same on any machine that can run the app.
"""

from __future__ import annotations

import csv
import io
import re

import pandas as pd
from docx import Document
from docx.text.paragraph import Paragraph
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

Section = tuple[str, list[dict[str, str]]]


def _columns(rows: list[dict[str, str]]) -> list[str]:
    return list(rows[0].keys()) if rows else []


def clean_rows(df: pd.DataFrame, required_columns: list[str] | None = None) -> list[dict[str, str]]:
    """Convert an edited ``st.data_editor`` DataFrame into clean
    string-valued rows for a :data:`Section`, dropping rows where every
    column in ``required_columns`` (default: all columns) is blank —
    covers the trailing empty row Streamlit's dynamic-rows editor always
    offers, plus any row a user added and then abandoned."""
    rows: list[dict[str, str]] = []
    for record in df.to_dict("records"):
        cleaned = {col: "" if pd.isna(value) else str(value) for col, value in record.items()}
        required = required_columns or list(cleaned.keys())
        if any(cleaned.get(col, "").strip() for col in required):
            rows.append(cleaned)
    return rows


def sections_to_docx(title: str, subtitle: str, sections: list[Section]) -> io.BytesIO:
    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        doc.add_paragraph(subtitle)
    for section_title, rows in sections:
        doc.add_heading(section_title, level=1)
        columns = _columns(rows)
        if not columns:
            doc.add_paragraph("None recorded.")
            continue
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        for cell, header in zip(table.rows[0].cells, columns, strict=False):
            cell.text = header
        for row in rows:
            cells = table.add_row().cells
            for cell, col in zip(cells, columns, strict=False):
                cell.text = str(row.get(col, "") or "")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def sections_to_xlsx(sections: list[Section]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    for section_title, rows in sections:
        name = (section_title or "Sheet")[:31]
        base, suffix_i = name, 1
        while name in used_names:
            suffix = f" ({suffix_i})"
            name = base[: 31 - len(suffix)] + suffix
            suffix_i += 1
        used_names.add(name)
        ws = wb.create_sheet(title=name)

        columns = _columns(rows)
        if not columns:
            ws.append(["None recorded."])
            continue
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(col, "") for col in columns])
        for idx, col in enumerate(columns, start=1):
            width = max([len(col)] + [len(str(row.get(col, "") or "")) for row in rows]) + 2
            ws.column_dimensions[get_column_letter(idx)].width = min(width, 60)

    if not wb.sheetnames:
        wb.create_sheet(title="Sheet1")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sections_to_csv(sections: list[Section]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for i, (section_title, rows) in enumerate(sections):
        if i > 0:
            writer.writerow([])
        writer.writerow([f"## {section_title}"])
        columns = _columns(rows)
        if not columns:
            writer.writerow(["None recorded."])
            continue
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(col, "") for col in columns])
    return buffer.getvalue()


# fpdf2's built-in core fonts (Helvetica etc.) only support Latin-1, but
# this app's own copy — dashes, curly quotes, arrows, middots — leans on
# Unicode punctuation throughout. Rather than bundle a Unicode TTF just
# for the PDF export, swap the common cases for ASCII equivalents and
# replace anything still outside Latin-1 (e.g. emoji in free-text fields)
# rather than letting fpdf2 raise.
_PDF_SAFE_REPLACEMENTS = {
    "—": "-",
    "–": "-",
    "‘": "'",
    "’": "'",
    "“": '"',
    "”": '"',
    "•": "-",
    "·": "-",
    "→": "->",
    "…": "...",
}


def _pdf_safe(text: str) -> str:
    for src, dst in _PDF_SAFE_REPLACEMENTS.items():
        text = text.replace(src, dst)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def sections_to_pdf(title: str, subtitle: str, sections: list[Section]) -> bytes:
    pdf = FPDF(orientation="L")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # multi_cell defaults to leaving the cursor at the end of the last
    # line rather than back at the left margin, which starves the next
    # multi_cell of horizontal space -- reset explicitly after each call.
    def _block(text: str, height: float) -> None:
        pdf.multi_cell(0, height, _pdf_safe(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font("Helvetica", "B", 16)
    _block(title, 10)
    if subtitle:
        pdf.set_font("Helvetica", "", 10)
        _block(subtitle, 6)
    pdf.ln(2)

    for section_title, rows in sections:
        pdf.set_font("Helvetica", "B", 13)
        _block(section_title, 8)
        pdf.set_font("Helvetica", "", 9)

        columns = _columns(rows)
        if not columns:
            _block("None recorded.", 6)
            pdf.ln(3)
            continue

        with pdf.table(text_align="LEFT") as table:
            header_row = table.row()
            for col in columns:
                header_row.cell(_pdf_safe(col))
            for row in rows:
                data_row = table.row()
                for col in columns:
                    data_row.cell(_pdf_safe(str(row.get(col, "") or "")))
        pdf.ln(3)

    return bytes(pdf.output())


# ---- Markdown documents (AI-drafted output, e.g. Discovery Call Assistant) ----
#
# Deliberately a small subset of Markdown — '#'/'##'/'###' headings, '-'/'*'
# and '1.' list items, **bold**, and plain paragraphs — which is exactly what
# the Discovery prompt asks the model to produce. Anything else (a stray
# table row, a code fence) falls through as a plain paragraph rather than
# failing the export.

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_BULLET = re.compile(r"^(\s*)[-*•]\s+(.*)$")
_MD_NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_MD_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _md_blocks(markdown: str) -> list[tuple[str, int, str]]:
    """``(kind, level, text)`` per line: kind is heading/bullet/numbered/para;
    level is the heading depth or the list indent depth."""
    blocks: list[tuple[str, int, str]] = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() in {"---", "***", "___"}:
            continue
        if m := _MD_HEADING.match(line.strip()):
            blocks.append(("heading", len(m.group(1)), m.group(2).strip().strip("*").strip()))
        elif m := _MD_BULLET.match(line):
            blocks.append(("bullet", len(m.group(1)) // 2, m.group(2)))
        elif m := _MD_NUMBERED.match(line):
            blocks.append(("numbered", len(m.group(1)) // 2, m.group(2)))
        else:
            blocks.append(("para", 0, line.strip()))
    return blocks


def _add_md_runs(paragraph: Paragraph, text: str) -> None:
    """Add ``text`` to a python-docx paragraph, rendering **bold** spans."""
    pos = 0
    for m in _MD_BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_to_docx(markdown: str, *, title: str = "") -> io.BytesIO:
    """Word document from Markdown. If the Markdown has no '#' title of its
    own, ``title`` is added as one."""
    doc = Document()
    blocks = _md_blocks(markdown)
    if title and not (blocks and blocks[0][0] == "heading" and blocks[0][1] == 1):
        doc.add_heading(title, level=0)
    for kind, level, text in blocks:
        if kind == "heading":
            doc.add_heading(text, level=0 if level == 1 else min(level - 1, 4))
            continue
        if kind in ("bullet", "numbered"):
            style = "List Bullet" if kind == "bullet" else "List Number"
            if level >= 1:
                style += " 2"
            paragraph = doc.add_paragraph(style=style)
        else:
            paragraph = doc.add_paragraph()
        _add_md_runs(paragraph, text)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def markdown_to_pdf(markdown: str, *, title: str = "") -> bytes:
    """Portrait PDF from Markdown, using fpdf2's own ``**bold**`` markup."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(18, 18, 18)
    pdf.add_page()

    def _write(text: str, height: float, *, indent: float = 0.0, markdown_text: bool = True) -> None:
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(
            pdf.epw - indent, height, _pdf_safe(text), markdown=markdown_text, new_x=XPos.LMARGIN, new_y=YPos.NEXT
        )

    blocks = _md_blocks(markdown)
    if title and not (blocks and blocks[0][0] == "heading" and blocks[0][1] == 1):
        blocks.insert(0, ("heading", 1, title))

    heading_sizes = {1: 18, 2: 14, 3: 12}
    numbers: dict[int, int] = {}
    for kind, level, text in blocks:
        if kind != "numbered":
            numbers.clear()
        if kind == "heading":
            pdf.ln(3 if level > 1 else 0)
            pdf.set_font("Helvetica", "B", heading_sizes.get(level, 11))
            _write(text, 8 if level <= 2 else 6, markdown_text=False)
            pdf.ln(1)
            continue
        pdf.set_font("Helvetica", "", 10)
        indent = 5.0 * (level + 1)
        if kind == "bullet":
            _write(f"-  {text}", 5.5, indent=indent)
        elif kind == "numbered":
            numbers[level] = numbers.get(level, 0) + 1
            _write(f"{numbers[level]}.  {text}", 5.5, indent=indent)
        else:
            _write(text, 5.5)
            pdf.ln(1.5)
    return bytes(pdf.output())
