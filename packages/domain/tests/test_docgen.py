"""Tests for Module 3 (Documentation Generator). Fully mocked — no live
portal or network access required.
"""

from docx import Document

from archscope_domain.docgen import DocumentationGenerator, snapshot_to_docx, snapshot_to_markdown
from archscope_integrations.hubspot import HubSpotScopeError


class FakeHubSpotClient:
    """Implements the subset of HubSpotClient's interface DocumentationGenerator
    calls, with canned data exercising every section of the doc pack."""

    def __init__(self, raise_on_companies: bool = False):
        self.raise_on_companies = raise_on_companies

    def get_properties(self, object_type: str) -> list[dict]:
        if object_type == "companies" and self.raise_on_companies:
            raise HubSpotScopeError("missing scope crm.schemas.companies.read")
        data = {
            "contacts": [
                {"name": "email", "label": "Email", "type": "string", "field_type": "text", "group_name": "contactinformation", "options": [], "description": "Contact email"},
                {"name": "lifecyclestage", "label": "Lifecycle Stage", "type": "enumeration", "field_type": "select", "group_name": "contactinformation", "options": [{"label": "Lead"}, {"label": "Customer"}], "description": ""},
            ],
            "companies": [
                {"name": "industry", "label": "Industry", "type": "enumeration", "field_type": "select", "group_name": "companyinformation", "options": [{"label": "Retail"}], "description": ""},
            ],
            "deals": [
                {"name": "amount", "label": "Amount", "type": "number", "field_type": "number", "group_name": "dealinformation", "options": [], "description": ""},
            ],
            "tickets": [
                {"name": "hs_pipeline_stage", "label": "Pipeline Stage", "type": "enumeration", "field_type": "select", "group_name": "ticketinformation", "options": [], "description": ""},
            ],
            "2-999": [
                {"name": "membership_status", "label": "Membership Status", "type": "enumeration", "field_type": "select", "group_name": "membershipinformation", "options": [{"label": "Active"}], "description": ""},
            ],
        }
        return data.get(object_type, [])

    def get_custom_object_schemas(self) -> list[dict]:
        return [
            {
                "object_type_id": "2-999",
                "name": "membership",
                "labels": {"plural": "Memberships", "singular": "Membership"},
                "associations": [
                    {"from_object_type_id": "2-999", "to_object_type_id": "0-1", "name": "membership_to_contact", "id": "1"}
                ],
            }
        ]

    def get_pipelines(self, object_type: str) -> list[dict]:
        data = {
            "deals": [
                {
                    "id": "default",
                    "label": "Sales Pipeline",
                    "stages": [
                        {"id": "s1", "label": "Appointment Scheduled", "display_order": 0},
                        {"id": "s2", "label": "Closed Won", "display_order": 1},
                    ],
                }
            ],
            "tickets": [],
            "2-999": [
                {"id": "m1", "label": "Membership Pipeline", "stages": [{"id": "ms1", "label": "New", "display_order": 0}]}
            ],
        }
        return data.get(object_type, [])

    def get_workflows(self) -> list[dict]:
        return [{"id": 111, "name": "Set Lifecycle Stage", "enabled": True, "type": "CONTACT", "reEnrollmentTriggersEnabled": True}]

    def get_workflow_detail(self, workflow_id) -> dict:
        return {
            "id": workflow_id,
            "name": "Set Lifecycle Stage",
            "enabled": True,
            "type": "CONTACT",
            "reEnrollmentTriggersEnabled": True,
            "actions": [{"type": "SET_PROPERTY", "propertyName": "lifecyclestage"}],
        }

    def get_owners(self) -> list[dict]:
        return [
            {
                "id": "1",
                "email": "jane@example.com",
                "first_name": "Jane",
                "last_name": "Doe",
                "teams": [{"name": "Sales", "id": "10", "primary": True}],
                "archived": False,
            }
        ]

    def get_teams(self) -> list[dict]:
        return [{"id": "10", "name": "Sales", "membership": ["1", "2"]}]


def test_build_snapshot_includes_standard_and_custom_objects():
    snapshot = DocumentationGenerator(FakeHubSpotClient()).build_snapshot()
    by_type = {s.object_type: s for s in snapshot.object_schemas}
    assert "contacts" in by_type
    assert by_type["contacts"].is_custom is False
    assert len(by_type["contacts"].properties) == 2
    assert "2-999" in by_type
    assert by_type["2-999"].is_custom is True
    assert by_type["2-999"].label == "Memberships"


def test_build_snapshot_pipelines_include_custom_object():
    snapshot = DocumentationGenerator(FakeHubSpotClient()).build_snapshot()
    by_object = {p.object_type: p for p in snapshot.pipelines}
    assert "Deal" in by_object
    assert [s.label for s in by_object["Deal"].stages] == ["Appointment Scheduled", "Closed Won"]
    assert "Memberships" in by_object


def test_build_snapshot_workflow_detail_merged():
    snapshot = DocumentationGenerator(FakeHubSpotClient()).build_snapshot()
    assert len(snapshot.workflows) == 1
    wf = snapshot.workflows[0]
    assert wf.name == "Set Lifecycle Stage"
    assert wf.re_enrollment_enabled is True
    assert len(wf.actions) == 1


def test_build_snapshot_owners_and_teams():
    snapshot = DocumentationGenerator(FakeHubSpotClient()).build_snapshot()
    assert snapshot.owners[0].email == "jane@example.com"
    assert snapshot.owners[0].teams == ["Sales"]
    assert snapshot.teams[0].name == "Sales"
    assert snapshot.teams[0].member_count == 2


def test_build_snapshot_association_map_from_custom_schema():
    snapshot = DocumentationGenerator(FakeHubSpotClient()).build_snapshot()
    membership_schema = next(s for s in snapshot.object_schemas if s.object_type == "2-999")
    assert len(membership_schema.associations) == 1
    assert membership_schema.associations[0].to_object == "0-1"


def test_build_snapshot_records_warning_without_blanking_other_sections():
    snapshot = DocumentationGenerator(FakeHubSpotClient(raise_on_companies=True)).build_snapshot()
    assert any("companies" in w.lower() for w in snapshot.warnings)
    by_type = {s.object_type: s for s in snapshot.object_schemas}
    assert by_type["companies"].properties == []
    # Other sections still populated despite the companies scope error.
    assert len(by_type["contacts"].properties) == 2
    assert len(snapshot.workflows) == 1


def test_snapshot_to_markdown_has_expected_sections():
    snapshot = DocumentationGenerator(FakeHubSpotClient()).build_snapshot()
    md = snapshot_to_markdown(snapshot)
    for heading in [
        "# HubSpot Portal Documentation",
        "## Data Dictionary",
        "## Workflow Inventory",
        "## Association Map",
        "## Pipelines & Stages",
        "## Roles & Permissions Summary",
    ]:
        assert heading in md
    assert "membership_status" in md
    assert "Sales Pipeline" in md


def test_snapshot_to_docx_produces_valid_document():
    snapshot = DocumentationGenerator(FakeHubSpotClient()).build_snapshot()
    buffer = snapshot_to_docx(snapshot)
    doc = Document(buffer)
    heading_texts = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    assert "HubSpot Portal Documentation" in heading_texts
    assert "Data Dictionary" in heading_texts
    assert "Workflow Inventory" in heading_texts
    assert "Association Map" in heading_texts
    assert "Pipelines & Stages" in heading_texts
    assert "Roles & Permissions Summary" in heading_texts
    assert len(doc.tables) > 0
