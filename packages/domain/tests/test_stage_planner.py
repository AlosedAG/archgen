"""Deal Stage Planner: the template round-trips, the parser reads whatever
the client did to the sheet (inserted / moved / relabelled rows, any
number of stage and field columns, placeholders, junk) without failing,
and the mappings feed Discovery, the Architecture Generator and the build
order."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from archscope_domain.blueprint import BlueprintGenerator, blueprint_to_markdown
from archscope_domain.discovery import discovery_build_order, empty_notes
from archscope_domain.models import BlueprintInput
from archscope_domain.rules import RulesEngine
from archscope_domain.stage_planner import (
    NOTES_END,
    NOTES_START,
    StagePlanner,
    StagePlannerConfigError,
    StagePlannerParseError,
    build_template,
    load_stage_planner_config,
    merge_notes,
    parse_stage_planner,
    parse_stage_planner_config,
    split_items,
    to_build_items,
    to_deal_properties,
    to_discovery_answer,
    to_discovery_notes,
    to_pipeline,
    to_workflows,
)


@pytest.fixture(scope="module")
def config():
    return load_stage_planner_config(discovery_build_order())


def _xlsx(rows: list[list[object]], *, title: str = "Planner", extra_sheets: tuple[str, ...] = ()) -> bytes:
    wb = Workbook()
    for name in extra_sheets:
        wb.create_sheet(name)
    ws = wb.active
    ws.title = title
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(row=r, column=c, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_round_trips_to_the_example_stages(config):
    planner = parse_stage_planner(build_template(config), config)
    assert planner.warnings == []
    assert [s.name for s in planner.stages] == [s["name"] for s in config.template["stages"]]
    won = next(s for s in planner.stages if s.name == "Closed Won")
    assert won.forecast == 100
    assert won.exit_criteria == ""  # "N/A — terminal stage" reads as blank
    assert won.fields == ["ARR/MRR", "Contract Start Date", "CSM Assigned", "Products Purchased"]


def test_rows_are_matched_by_label_not_position(config):
    content = _xlsx(
        [
            ["My pipeline"],
            [],
            ["  stage name: ", "Discovery", "Proposal", "Won deal"],
            ["Notes row the client added", "x", "y", "z"],
            ["Automations", "Notify owner", "", "Create onboarding task"],
            ["Probability", "15%", "50%", "100%"],
            ["Exit criteria", "Call booked", "Signed", "N/A"],
            ["Fields to fill out"],
            ["Field 1", "Lead Source", "Amount", "ARR"],
            [None, "Budget", None, None],  # unlabelled field row with a value
            ["Field 9", None, "Lead Source", None],
        ]
    )
    planner = parse_stage_planner(content, config)
    assert [s.name for s in planner.stages] == ["Discovery", "Proposal", "Won deal"]
    assert [s.forecast for s in planner.stages] == [15, 50, 100]
    assert planner.stages[0].exit_criteria == "Call booked"
    assert planner.stages[2].exit_criteria == ""
    assert planner.stages[0].fields == ["Lead Source", "Budget"]
    assert planner.stages[1].fields == ["Amount", "Lead Source"]
    assert any("'Entry Criteria'" in w for w in planner.warnings)
    assert any("'Desired Actions / Activities'" in w for w in planner.warnings)


def test_blank_placeholder_and_unnamed_columns_are_skipped_with_warnings(config):
    content = _xlsx(
        [
            ["Stage Name", "Lead", None, "Stage name", "Stage name", "Deal"],
            ["Forecast %", 10, 30, None, 40, 90],
            ["Fields to Fill Out"],
            ["Field 1", "Source", "Orphan field", None, None, None],
        ]
    )
    planner = parse_stage_planner(content, config)
    assert [s.name for s in planner.stages] == ["Lead", "Deal"]
    assert any("Column C has details but no stage name" in w for w in planner.warnings)
    assert any("Column E still says 'Stage name'" in w for w in planner.warnings)
    assert not any("Column D" in w for w in planner.warnings)  # empty placeholder column: silently skipped


def test_forecasts_as_fractions_and_bad_values(config):
    fractions = parse_stage_planner(_xlsx([["Stage Name", "A", "B"], ["Forecast %", 0.25, 1]]), config)
    assert [s.forecast for s in fractions.stages] == [25, 100]
    assert any("fractions" in w for w in fractions.warnings)
    bad = parse_stage_planner(_xlsx([["Stage Name", "A", "B", "C"], ["Forecast %", "high", 140, " 35 % "]]), config)
    assert [s.forecast for s in bad.stages] == [None, None, 35]
    assert any("'high' isn't a number" in w for w in bad.warnings)
    assert any("outside 0–100%" in w for w in bad.warnings)


def test_whitespace_duplicates_and_sheet_fallback(config):
    content = _xlsx(
        [["Deal Stage", "  New \n  lead  ", "Won", "won"], ["Entry Criteria", "  first   touch  "]],
        title="Sheet1",
        extra_sheets=("Instructions",),
    )
    planner = parse_stage_planner(content, config)
    assert planner.source == "Sheet1"
    assert planner.stages[0].name == "New\nlead"
    assert planner.stages[0].entry_criteria == "first touch"
    assert any("appears twice" in w for w in planner.warnings)


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"hello", "isn't an .xlsx workbook"),
        (_xlsx([["Something else", "x"]]), "doesn't look like a Stage Planner"),
    ],
)
def test_unreadable_files_raise(config, content, message):
    with pytest.raises(StagePlannerParseError, match=message):
        parse_stage_planner(content, config)


def test_no_stages_or_fields_section_are_warnings_not_errors(config):
    planner = parse_stage_planner(_xlsx([["Stage Name", "Stage name"]]), config)
    assert planner.stages == []
    assert any("No stages found" in w for w in planner.warnings)
    assert any("No 'Fields to Fill Out' section" in w for w in planner.warnings)


# ---- Mapping ----------------------------------------------------------------


@pytest.fixture(scope="module")
def sample(config):
    return parse_stage_planner(build_template(config), config)


def test_pipeline_uses_order_probability_and_criteria(config, sample):
    pipeline = to_pipeline(sample, config)
    assert pipeline.object_type == "Deal"
    stages = {s.name: s for s in pipeline.stages}
    assert stages["Qualified to Buy"].probability == 0.2
    assert stages["Qualified to Buy"].description.startswith("Entry: Intro call completed")
    assert (stages["Closed Won"].is_closed, stages["Closed Won"].is_won) == (True, True)
    assert (stages["Closed Lost"].is_closed, stages["Closed Lost"].is_won) == (True, False)
    assert [s.order for s in pipeline.stages] == list(range(len(sample.stages)))


def test_deal_properties_are_typed_tagged_and_mapped_to_existing(config, sample):
    props = {p.label: p for p in to_deal_properties(sample, config)}
    assert props["Demo Date"].definition.type == "date"
    assert props["Discount %"].definition.name == "discount_pct"
    assert props["Discount %"].definition.type == "number"
    assert props["Lead Source"].definition.description == (
        "Required at stage: Appointment Scheduled (from the Deal Stage Planner)."
    )
    assert props["Expected Close Date"].existing == "closedate"
    assert props["Expected Close Date"].definition is None


def test_a_field_required_at_several_stages_is_one_property(config):
    planner = StagePlanner.from_dict(
        {"stages": [{"name": "A", "fields": ["Budget"]}, {"name": "B", "fields": ["budget ", "Champion"]}]}
    )
    props = to_deal_properties(planner, config)
    assert [(p.label, p.stages) for p in props] == [("Budget", ("A", "B")), ("Champion", ("B",))]


def test_workflows_and_build_items_from_automations_and_actions(config, sample):
    workflows = to_workflows(sample, config, RulesEngine())
    assert workflows[0].name == "Deal | Appointment Scheduled: Create task: send intro email"
    assert workflows[0].trigger == "Deal enters the 'Appointment Scheduled' stage"
    items = to_build_items(sample, config)
    assert items[0].section == "pipelines"
    assert any(i.name == "Rep task / sequence — Contract Sent: Loop in finance and legal" for i in items)
    assert not any(i.name == "Deal property: Expected Close Date" for i in items)  # mapped to closedate
    assert len({i.key for i in items}) == len(items)


def test_build_order_includes_planner_items(config, sample):
    order = discovery_build_order().build(empty_notes(), extra_items=to_build_items(sample, config))
    pipelines = dict((s.key, items) for s, items in order.sections)["pipelines"]
    assert pipelines[0].name == "Deal pipeline from the Stage Planner (7 stages)"


def test_blueprint_uses_the_imported_pipeline_properties_and_workflows(config, sample):
    pipeline = to_pipeline(sample, config)
    props = [p.definition for p in to_deal_properties(sample, config) if p.definition]
    blueprint = BlueprintGenerator(RulesEngine()).generate(
        BlueprintInput(
            project_name="Acme",
            standard_objects=["Deal"],
            pipeline_overrides={"Deal": pipeline.stages},
            additional_properties={"Deal": props},
            additional_workflows=to_workflows(sample, config),
        )
    )
    (deal,) = blueprint.pipelines
    assert [s.name for s in deal.stages] == [s.name for s in sample.stages]
    assert any(p.name == "demo_date" for p in blueprint.standard_object_properties["Deal"])
    assert any(w.trigger.startswith("Deal enters") for w in blueprint.workflows)
    markdown = blueprint_to_markdown(blueprint)
    assert "| Order | Stage | Probability | Closed | Won | Description |" in markdown
    assert "| 1 | Qualified to Buy | 20% |" in markdown


def test_discovery_answer_notes_and_merge(config, sample):
    assert to_discovery_answer(sample).startswith("Appointment Scheduled (10%) → Qualified to Buy (20%)")
    block = to_discovery_notes(sample, config)
    assert block.startswith(NOTES_START)
    assert block.endswith(NOTES_END)
    assert "  Required fields: Lead Source, Company Size, Initial Notes" in block
    merged = merge_notes("Call notes here.", block)
    assert merged.startswith("Call notes here.\n\n" + NOTES_START)
    # A re-import replaces the earlier block rather than adding another.
    again = merge_notes(merged, f"{NOTES_START}\nNew\n{NOTES_END}")
    assert again.count(NOTES_START) == 1
    assert "New" in again
    assert again.startswith("Call notes here.")


def test_split_items(config):
    assert split_items("Notify AE; schedule demo\n• log feedback.", config) == ["Notify AE", "Schedule demo", "Log feedback"]


def test_planner_round_trips_through_dict(sample):
    assert StagePlanner.from_dict(sample.to_dict()) == sample


# ---- Config validation --------------------------------------------------------


def test_config_problems_are_reported_together():
    build = discovery_build_order()
    raw = {
        "rows": [{"key": "forecast", "label": "Forecast %"}, {"key": "nope", "label": "X"}],
        "fields_section": {"label": "Fields", "row_label": "Field"},
        "template": {"title": "T", "stages": [{"name": "A", "forecast": "lots"}]},
        "mapping": {
            "discovery": {"section": "Processes", "question_id": "journey_stages"},
            "build_order": {"pipeline": {"section": "nowhere", "name": "x"}},
        },
    }
    with pytest.raises(StagePlannerConfigError) as exc:
        parse_stage_planner_config(raw, build)
    problems = "\n".join(exc.value.problems)
    for expected in [
        "rows #2",
        "the 'name' row (stage names) is required",
        "must contain {n}",
        "'forecast' must be a number",
        "isn't a text question in Processes",
        "build_order › pipeline",
    ]:
        assert expected in problems, problems
