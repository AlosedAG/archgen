"""Tests for Module 4 (Executive Report). Fully offline — builds fixtures
directly, no HubSpot client or network access involved.
"""

from docx import Document

from archscope_domain.models import (
    AssociationDef,
    Finding,
    ObjectSchema,
    Owner,
    PortalPipeline,
    PortalPipelineStage,
    PortalPropertyDef,
    PortalSnapshot,
    PortalWorkflow,
    Team,
)
from archscope_domain.report import (
    build_report_context,
    prettify_group_name,
    report_to_client_docx,
    report_to_internal_docx,
)


def _all_text(doc: Document) -> str:
    """Flatten every paragraph and table cell into one searchable string."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _snapshot() -> PortalSnapshot:
    contacts = ObjectSchema(
        object_type="contacts",
        label="Contact",
        is_custom=False,
        properties=[
            PortalPropertyDef(name="email", label="Email", type="string", group_name="contactinformation", hubspot_defined=True),
            PortalPropertyDef(
                name="lifecyclestage",
                label="Lifecycle Stage",
                type="enumeration",
                group_name="contactinformation",
                hubspot_defined=True,
            ),
        ],
    )
    locations = ObjectSchema(
        object_type="2-999",
        label="Locations",
        is_custom=True,
        singular_label="Location",
        properties=[
            PortalPropertyDef(name="location_name", label="Location Name", type="string", group_name="locations_information"),
        ],
        associations=[
            AssociationDef(from_object="Location", to_object="0-1", label="location_to_contact"),
            AssociationDef(from_object="Location", to_object="2-999", label="location_to_location"),
        ],
    )
    pipeline = PortalPipeline(
        object_type="Deal",
        pipeline_id="p1",
        label="Sales Pipeline",
        stages=[
            PortalPipelineStage(stage_id="s1", label="New", display_order=0),
            PortalPipelineStage(stage_id="s2", label="Won", display_order=1),
        ],
    )
    workflows = [
        PortalWorkflow(workflow_id="w1", name="Contact | Welcome Email", enabled=True, object_type="Contact"),
        PortalWorkflow(workflow_id="w2", name="Unnamed workflow - Mon Jan 1 2026", enabled=False, object_type="Contact"),
    ]
    owner_unassigned = Owner(owner_id="1", email="a@x.com", first_name="A", last_name="B", teams=[], archived=False)
    owner_on_team = Owner(owner_id="2", email="c@x.com", first_name="C", last_name="D", teams=["Sales"], archived=False)
    team_empty = Team(team_id="10", name="Empty Team", member_count=0)
    team_staffed = Team(team_id="11", name="Sales", member_count=3)

    return PortalSnapshot(
        pulled_at="2026-01-01T00:00:00+00:00",
        object_schemas=[contacts, locations],
        pipelines=[pipeline],
        workflows=workflows,
        owners=[owner_unassigned, owner_on_team],
        teams=[team_empty, team_staffed],
    )


def _findings() -> list[Finding]:
    return [
        Finding(
            area="Workflows",
            object_type="Contact",
            description="Workflow 'X' sets a property unconditionally.",
            severity="High",
            recommended_fix="Add an 'only if blank' branch.",
        ),
        Finding(
            area="Permissions",
            object_type="Team",
            description="Team 'Empty Team' has zero members.",
            severity="Medium",
            recommended_fix="Confirm the team is still needed.",
        ),
        Finding(
            area="Naming",
            object_type="Contact",
            description="Property 'Lead_Score' naming drift.",
            severity="Low",
            recommended_fix="Rename to snake_case.",
        ),
        Finding(
            area="Naming",
            object_type="Contact",
            description="Property 'Other_Prop' naming drift.",
            severity="Low",
            recommended_fix="Rename to snake_case.",
        ),
    ]


def test_prettify_group_name_handles_known_and_generic():
    assert prettify_group_name("contactinformation") == "Contact information"
    assert prettify_group_name("locations_information") == "Locations information"
    assert prettify_group_name("") == "Ungrouped"


def test_build_report_context_computes_object_and_pipeline_stats():
    ctx = build_report_context(_snapshot(), _findings(), project_name="Acme")
    assert ctx.project_name == "Acme"
    assert ctx.standard_object_count == 1
    assert ctx.custom_object_count == 1
    assert ctx.total_properties == 3
    assert ctx.total_pipelines == 1
    assert ctx.total_stages == 2

    locations = next(o for o in ctx.object_summaries if o.label == "Locations")
    assert locations.association_count == 2
    assert "Contact" in locations.association_targets
    assert "Locations" in locations.association_targets  # self-association resolved via snapshot's own label


def test_build_report_context_workflow_and_permission_stats():
    ctx = build_report_context(_snapshot(), _findings())
    assert ctx.workflow_total == 2
    assert ctx.workflow_enabled == 1
    assert ctx.workflow_disabled == 1
    assert ctx.workflow_unnamed == 1
    assert ctx.owner_unassigned == 1
    assert ctx.team_empty == 1


def test_build_report_context_groups_findings_and_rates_health():
    ctx = build_report_context(_snapshot(), _findings())
    assert ctx.audited is True
    assert ctx.finding_total == 4
    assert ctx.severity_counts == {"High": 1, "Medium": 1, "Low": 2}
    assert ctx.rating == "Needs Attention"
    # High-severity group sorts first.
    assert ctx.finding_groups[0].severity == "High"
    naming_low_group = next(g for g in ctx.finding_groups if g.area == "Naming" and g.severity == "Low")
    assert naming_low_group.count == 2


def test_build_report_context_without_audit_marks_not_audited():
    ctx = build_report_context(_snapshot(), None)
    assert ctx.audited is False
    assert ctx.finding_total == 0
    assert ctx.rating == "Not Yet Audited"
    assert ctx.finding_groups == []


def test_build_report_context_clean_audit_rates_excellent():
    ctx = build_report_context(_snapshot(), [])
    assert ctx.audited is True
    assert ctx.rating == "Excellent"


def test_report_to_client_docx_produces_valid_document_with_expected_sections():
    ctx = build_report_context(_snapshot(), _findings(), project_name="Acme")
    buffer = report_to_client_docx(ctx)
    doc = Document(buffer)
    heading_texts = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    assert any("Acme" in h for h in heading_texts)
    for expected in [
        "Overall health",
        "What's in your portal",
        "Automation",
        "Users & teams",
        "Things worth fixing",
        "Recommended next steps",
    ]:
        assert expected in heading_texts
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Needs Attention" in body_text


def test_report_to_client_docx_caps_examples_per_group():
    many_findings = [
        Finding(
            area="Naming",
            object_type="Contact",
            description=f"Property 'p{i}' naming drift.",
            severity="Low",
            recommended_fix="Rename.",
        )
        for i in range(10)
    ]
    ctx = build_report_context(_snapshot(), many_findings)
    buffer = report_to_client_docx(ctx)
    doc = Document(buffer)
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "...and 7 more like this." in body_text


def test_report_to_internal_docx_includes_full_finding_detail():
    ctx = build_report_context(_snapshot(), _findings(), project_name="Acme")
    buffer = report_to_internal_docx(ctx)
    doc = Document(buffer)
    heading_texts = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    assert any("Acme" in h for h in heading_texts)
    for expected in ["Health rating", "Portal stats", "Object model", "Findings — full detail", "Summary by area and priority"]:
        assert expected in heading_texts

    text = _all_text(doc)
    # Both Naming/Low findings appear individually (no capping in the internal report).
    assert "Lead_Score" in text
    assert "Other_Prop" in text
    assert "locations_information" not in text  # raw group id shouldn't leak; prettified name should
    assert "Locations information" in text
    # Findings live in tables (not just paragraphs).
    assert len(doc.tables) >= 3


def test_report_to_internal_docx_without_audit_notes_it():
    ctx = build_report_context(_snapshot(), None)
    buffer = report_to_internal_docx(ctx)
    doc = Document(buffer)
    assert "No audit has been run yet" in _all_text(doc)
