"""Tests for Module 2B (Property Audit). Fully mocked — no live portal or
network access required.
"""

from __future__ import annotations

import io

import openpyxl
from pypdf import PdfReader

from core.hubspot_client import HubSpotScopeError
from core.models import ObjectSchema, PortalPropertyDef, PortalSnapshot, PortalWorkflow
from core.property_audit import (
    RATING_KEEP,
    RATING_REMOVE,
    RATING_REVIEW,
    _rate_property,
    property_audit_to_pdf,
    property_audit_to_xlsx,
    run_property_audit,
)


class FakeHubSpotClient:
    """Returns a canned, fixed-fill-rate sample per object type so rating
    boundaries are deterministic rather than randomized."""

    def __init__(self, fill_fraction_by_property: dict[str, float], sample_size: int = 100):
        self.fill_fraction_by_property = fill_fraction_by_property
        self.sample_size = sample_size

    def get_records_sample(self, object_type, limit=100, properties=None):
        n = min(limit, self.sample_size)
        records = []
        for i in range(n):
            props = {}
            for name in properties or []:
                fraction = self.fill_fraction_by_property.get(name, 0.0)
                if i < round(fraction * n):
                    props[name] = "x"
            records.append({"id": str(i), "properties": props})
        return records


# ---- rating boundary logic (reverse-engineered from examples/RPG_property_audit.xlsx) ----


def test_native_property_is_always_keep_regardless_of_fill_or_uses():
    rating, assessment = _rate_property(is_custom=False, fill_pct=0.0, uses=0)
    assert rating == RATING_KEEP
    assert assessment == ""


def test_custom_unused_under_five_percent_fill_is_remove_candidate():
    rating, _ = _rate_property(is_custom=True, fill_pct=4.9, uses=0)
    assert rating == RATING_REMOVE


def test_custom_unused_at_five_percent_fill_is_review_not_remove():
    rating, assessment = _rate_property(is_custom=True, fill_pct=5.2, uses=0)
    assert rating == RATING_REVIEW
    assert "likely removable" in assessment


def test_custom_unused_fully_filled_is_review_verify_before_removing():
    rating, assessment = _rate_property(is_custom=True, fill_pct=100.0, uses=0)
    assert rating == RATING_REVIEW
    assert "verify before removing" in assessment


def test_custom_used_under_25_percent_fill_is_review():
    rating, assessment = _rate_property(is_custom=True, fill_pct=24.9, uses=1)
    assert rating == RATING_REVIEW
    assert "review" in assessment.lower()


def test_custom_used_at_25_percent_fill_is_keep():
    rating, assessment = _rate_property(is_custom=True, fill_pct=25.2, uses=1)
    assert rating == RATING_KEEP
    assert assessment == ""


# ---- run_property_audit end-to-end ----------------------------------------


def _snapshot() -> PortalSnapshot:
    contacts = ObjectSchema(
        object_type="contacts",
        label="Contact",
        is_custom=False,
        properties=[
            PortalPropertyDef(name="email", label="Email", type="string", hubspot_defined=True),
            PortalPropertyDef(name="dead_prop", label="Dead Prop", type="string", hubspot_defined=False),
            PortalPropertyDef(name="referenced_prop", label="Referenced Prop", type="string", hubspot_defined=False),
        ],
    )
    wf = PortalWorkflow(
        workflow_id="1",
        name="Contact | Set referenced_prop",
        enabled=True,
        object_type="CONTACT",
        actions=[{"type": "SET_PROPERTY", "propertyName": "referenced_prop"}],
    )
    return PortalSnapshot(pulled_at="2026-01-01T00:00:00+00:00", object_schemas=[contacts], workflows=[wf])


def test_run_property_audit_rates_dead_and_referenced_properties_correctly():
    client = FakeHubSpotClient({"email": 1.0, "dead_prop": 0.0, "referenced_prop": 0.1})
    result = run_property_audit(client, _snapshot())

    by_name = {r.internal_name: r for r in result.rows}
    assert by_name["email"].rating == RATING_KEEP  # native
    assert by_name["dead_prop"].rating == RATING_REMOVE
    assert by_name["dead_prop"].uses == 0
    assert by_name["referenced_prop"].uses == 1
    assert by_name["referenced_prop"].rating == RATING_REVIEW

    assert len(result.object_summaries) == 1
    summary = result.object_summaries[0]
    assert summary.properties_count == 3
    assert summary.custom_count == 2
    assert summary.remove_count == 1
    assert summary.review_count == 1
    assert summary.keep_count == 1
    # cleanup_score = 100 * (1 - (1 + 0.5*1) / 2) = 25
    assert summary.cleanup_score == 25

    assert result.totals.cleanup_score == 25
    assert {r.internal_name for r in result.flagged_rows} == {"dead_prop", "referenced_prop"}


def test_run_property_audit_records_warning_on_scope_error_but_does_not_raise():
    class FailingClient:
        def get_records_sample(self, object_type, limit=100, properties=None):
            raise HubSpotScopeError("missing scope")

    result = run_property_audit(FailingClient(), _snapshot())
    assert result.warnings
    assert all(r.fill_pct == 0.0 for r in result.rows)


# ---- exporters ---------------------------------------------------------------


def test_property_audit_to_xlsx_has_expected_sheets_and_header_row():
    client = FakeHubSpotClient({"email": 1.0, "dead_prop": 0.0, "referenced_prop": 0.1})
    result = run_property_audit(client, _snapshot())

    wb = openpyxl.load_workbook(io.BytesIO(property_audit_to_xlsx(result, project_name="Acme")))
    assert wb.sheetnames[:4] == ["Summary", "All Properties", "Flagged for Action", "How to use"]
    all_props_ws = wb["All Properties"]
    assert [c.value for c in all_props_ws[1]] == [
        "Object",
        "Property",
        "Internal name",
        "Type",
        "Custom",
        "Fill %",
        "Uses",
        "Rating",
        "Assessment",
        "Notes",
        "Tag",
    ]
    assert all_props_ws.max_row == len(result.rows) + 1


def test_property_audit_to_pdf_stays_within_ten_pages_for_a_large_portal():
    fill_by_name = {"native": 1.0}
    schemas = []
    for oi in range(30):
        props = [PortalPropertyDef(name="native", label="Native", type="string", hubspot_defined=True)]
        for i in range(10):
            name = f"obj{oi}_custom_{i}"
            props.append(PortalPropertyDef(name=name, label=f"Custom {i}", type="string", hubspot_defined=False))
            fill_by_name[name] = 0.0
        schemas.append(ObjectSchema(object_type=f"obj_{oi}", label=f"Object {oi}", is_custom=True, properties=props))
    snapshot = PortalSnapshot(pulled_at="2026-01-01T00:00:00+00:00", object_schemas=schemas)

    client = FakeHubSpotClient(fill_by_name)
    result = run_property_audit(client, snapshot)

    pdf_bytes = property_audit_to_pdf(result, project_name="BigCo")
    assert len(PdfReader(io.BytesIO(pdf_bytes)).pages) <= 10


def test_property_audit_to_pdf_handles_empty_result():
    result = run_property_audit(FakeHubSpotClient({}), PortalSnapshot(pulled_at="2026-01-01T00:00:00+00:00"))
    pdf_bytes = property_audit_to_pdf(result, project_name="Empty")
    assert len(PdfReader(io.BytesIO(pdf_bytes)).pages) >= 1
