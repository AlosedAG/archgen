import csv
import io

import numpy as np
import openpyxl
import pandas as pd
from docx import Document
from pypdf import PdfReader

from core.exporters import clean_rows, sections_to_csv, sections_to_docx, sections_to_pdf, sections_to_xlsx

SECTIONS = [
    ("Customer Requirements", [{"Requirement": "Sync contacts nightly"}, {"Requirement": "Log SMS replies"}]),
    ("Open Questions", []),
]


def test_sections_to_docx_includes_headings_and_table_rows():
    buffer = sections_to_docx("Test WRD", "A subtitle", SECTIONS)
    doc = Document(buffer)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Test WRD" in text
    assert "A subtitle" in text
    assert "Customer Requirements" in text
    assert "Open Questions" in text
    assert "None recorded." in text

    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "Requirement"
    assert table.rows[1].cells[0].text == "Sync contacts nightly"


def test_sections_to_xlsx_creates_one_sheet_per_section():
    data = sections_to_xlsx(SECTIONS)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Customer Requirements" in wb.sheetnames
    assert "Open Questions" in wb.sheetnames
    ws = wb["Customer Requirements"]
    assert ws["A1"].value == "Requirement"
    assert ws["A2"].value == "Sync contacts nightly"
    assert wb["Open Questions"]["A1"].value == "None recorded."


def test_sections_to_xlsx_dedupes_long_and_duplicate_sheet_names():
    long_name = "A" * 40
    data = sections_to_xlsx([(long_name, [{"X": "1"}]), (long_name, [{"X": "2"}])])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert len(wb.sheetnames) == 2
    assert all(len(name) <= 31 for name in wb.sheetnames)
    assert len(set(wb.sheetnames)) == 2


def test_sections_to_csv_has_section_markers_and_rows():
    text = sections_to_csv(SECTIONS)
    reader = list(csv.reader(io.StringIO(text)))
    assert ["## Customer Requirements"] in reader
    assert ["Requirement"] in reader
    assert ["Sync contacts nightly"] in reader
    assert ["## Open Questions"] in reader
    assert ["None recorded."] in reader


def test_sections_to_pdf_produces_valid_pdf_with_expected_text():
    data = sections_to_pdf("Test WRD", "A subtitle", SECTIONS)
    assert data[:4] == b"%PDF"
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Test WRD" in text
    assert "Customer Requirements" in text
    assert "Sync contacts nightly" in text
    assert "None recorded." in text


def test_sections_to_pdf_handles_empty_sections_list():
    data = sections_to_pdf("Empty Doc", "", [])
    assert data[:4] == b"%PDF"


def test_sections_to_pdf_survives_unicode_punctuation_used_elsewhere_in_the_app():
    # fpdf2's core Helvetica font is Latin-1 only; this app's own copy uses
    # em dashes, arrows, and middots throughout (see e.g. page headers and
    # diagram captions), so the PDF exporter must not choke on them.
    data = sections_to_pdf(
        "Written Requirements Document — Acme Corp",
        "Documents the customer's business objectives — see appendix for the ERD.",
        [("Findings — Overview", [{"Object": "Contact -> Deal", "Note": "3 findings · High priority"}])],
    )
    assert data[:4] == b"%PDF"
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Acme Corp" in text


def test_clean_rows_drops_blank_trailing_rows_from_dynamic_editor():
    df = pd.DataFrame(
        [
            {"Requirement": "Sync contacts nightly"},
            {"Requirement": "  "},
            {"Requirement": None},
        ]
    )
    assert clean_rows(df) == [{"Requirement": "Sync contacts nightly"}]


def test_clean_rows_converts_nan_and_none_to_empty_string():
    df = pd.DataFrame([{"Term": "SKU", "Definition": np.nan}, {"Term": None, "Definition": None}])
    assert clean_rows(df, required_columns=["Term"]) == [{"Term": "SKU", "Definition": ""}]


def test_clean_rows_keeps_row_when_any_required_column_has_content():
    df = pd.DataFrame([{"Type": "Risk", "Description": ""}])
    assert clean_rows(df, required_columns=["Type", "Description"]) == [{"Type": "Risk", "Description": ""}]


def test_markdown_document_exports_round_trip():
    import io as _io

    from docx import Document as _Document
    from pypdf import PdfReader

    from core.exporters import markdown_to_docx, markdown_to_pdf

    md = "# Discovery — Acme\n## Executive Summary\nA **dental** group → growth.\n- one\n  - nested\n1. first\n"
    doc = _Document(markdown_to_docx(md, title="ignored, markdown has its own title"))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Discovery — Acme"
    assert "A dental group → growth." in texts
    assert any(r.bold and r.text == "dental" for p in doc.paragraphs for r in p.runs)

    pdf_text = PdfReader(_io.BytesIO(markdown_to_pdf(md))).pages[0].extract_text()
    assert "Executive Summary" in pdf_text and "growth" in pdf_text


def test_markdown_export_adds_title_when_missing():
    from docx import Document as _Document

    from core.exporters import markdown_to_docx

    doc = _Document(markdown_to_docx("## Section\nBody", title="My Title"))
    assert doc.paragraphs[0].text == "My Title"
